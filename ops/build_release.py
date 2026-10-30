from pathlib import Path
import sys, shutil, subprocess, zipfile, hashlib, json
from pypdf import PdfReader,PdfWriter,Transformation
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor,white
from reportlab.lib.utils import simpleSplit
from openpyxl import load_workbook
src=Path(sys.argv[1]);out=Path(sys.argv[2]);out.mkdir(parents=True,exist_ok=True)
binder=src/'approved-binder-v8.pdf';workbook=src/'calculator-v2.xlsx'
r=PdfReader(binder);assert len(r.pages)==64 and not r.get_fields()
w=load_workbook(workbook,data_only=False);cached=load_workbook(workbook,data_only=True);assert len(w.sheetnames)==9
for sheet in w:
 for row in sheet:
  for cell in row:
   if cell.data_type=='f':
    value=cached[sheet.title][cell.coordinate].value
    assert value is not None and value not in ['#REF!','#VALUE!','#DIV/0!','#NAME?','#NUM!'],(sheet.title,cell.coordinate)
shutil.copyfile(binder,out/'Household-Binder-US-Letter.pdf');shutil.copyfile(workbook,out/'Household-Calculator.xlsx')
a4=(595.2756,841.8898)
pw=PdfWriter()
for p in PdfReader(binder).pages:
 blank=pw.add_blank_page(*a4);scale=min(a4[0]/float(p.mediabox.width),a4[1]/float(p.mediabox.height));blank.merge_transformed_page(p,Transformation().scale(scale).translate((a4[0]-float(p.mediabox.width)*scale)/2,(a4[1]-float(p.mediabox.height)*scale)/2))
with open(out/'Household-Binder-A4.pdf','wb') as f:pw.write(f)
for paper in ['US-Letter','A4']:
 subprocess.run(['gs','-q','-dBATCH','-dNOPAUSE','-sDEVICE=pdfwrite','-sColorConversionStrategy=Gray','-dProcessColorModel=/DeviceGray','-sOutputFile='+str(out/f'Household-Binder-{paper}-Grayscale.pdf'),str(out/f'Household-Binder-{paper}.pdf')],check=True)
green=HexColor('#173d2e');ink=HexColor('#172119')
def text(c,t,x,y,width=510,size=11):
 c.setFillColor(ink);c.setFont('Helvetica',size)
 for line in simpleSplit(t,'Helvetica',size,width):c.drawString(x,y,line);y-=size*1.45
 return y
c=canvas.Canvas(str(out/'Start-Here.pdf'),pagesize=(612,792));c.setTitle('LIVSOLV Systems Down Household Binder - Start Here');c.setFillColor(green);c.setFont('Helvetica-Bold',23);c.drawString(48,735,'START HERE');y=700
for title,body in [
('1. Choose your print edition','US Letter fits standard U.S. paper. A4 fits A4 paper. Grayscale editions use black-and-white colors; large dark areas may still use ink. Print one edition, not all four. The binder has 62 numbered pages plus its cover and Emergency Quick Find.'),
('2. Start with 15 minutes','Open numbered page 2 for the basic setup. Fill in the inside-front-cover Emergency Quick Find, household contacts and essential supply locations first. This PDF is designed to print and write on; it is not an interactive fillable form.'),
('3. Open the calculator','Use Household-Calculator.xlsx in a spreadsheet app that supports Excel formulas. Read its Start Here tab before entering household information. Change input cells, preserve formulas, and save your own copy. Figures are planning estimates, not equipment guarantees.'),
('4. Place the Quick Cards','Print Quick-Cards.pdf, cut along the borders, and fill in the household-specific details. Keep relevant cards with supplies or in a wallet. Avoid writing full passwords, account numbers or identity-document numbers.'),
('5. Make the plan easy to find','Tell household members where the printed binder lives. Use the tiered jobs page to assign primary and backup roles. Print extra copies of the expense log and section-notes sheet as needed.'),
('If an emergency is happening now','Open numbered page 1, A Major Emergency Just Happened. Follow local emergency instructions and prioritize immediate safety. Do not delay leaving danger to complete paperwork.')]:
 c.setFillColor(green);c.setFont('Helvetica-Bold',12);c.drawString(48,y,title);y-=20;y=text(c,body,48,y);y-=19
c.setFont('Helvetica',9);c.drawString(48,35,'LIVSOLV.com | Live in the solution.');c.save()
cards=[('HOUSEHOLD CONTACTS',['Out-of-area contact','Phone / messaging app','Household meeting place','Backup meeting place']),('EMERGENCY QUICK FIND',['Flashlight / headlamp location','Charged power bank location','First-aid kit location','Binder location']),('LEAVE-HOME PLAN',['Primary destination','Backup destination','Route / local map location','Person responsible for go bags']),('HOUSEHOLD JOBS',['Immediate safety lead','People / pet check lead','Supplies / communications lead','Backup person']),('WATER PLAN',['Drinking-water supply location','Household daily target','Backup water source','Treatment instructions location']),('POWER PLAN',['Essential device / backup','Chargers / cables location','Backup power location','Device instructions location']),('FOOD PLAN',['Eat-first plan location','Shelf-stable foods location','Manual can opener location','Thermometer location']),('MEDICAL CONTINUITY',['Critical medication location','Device / power backup','Care-team contact','Protected medical plan location']),('PETS AND DEPENDENTS',['Person responsible','Food / medication location','Carrier / leash location','Backup caregiver']),('COMMUNICATIONS',['Out-of-area check-in contact','Check-in time / method','Wi-Fi availability / backup','Radio / alerts source'])]
c=canvas.Canvas(str(out/'Quick-Cards.pdf'),pagesize=(612,792));c.setTitle('LIVSOLV Household Emergency Quick Cards')
for i,(title,fields) in enumerate(cards):
 if i and i%2==0:c.showPage()
 top=744-(i%2)*354;c.setStrokeColor(green);c.rect(42,top-314,528,314);c.setFillColor(green);c.rect(42,top-43,528,43,fill=1,stroke=0);c.setFillColor(white);c.setFont('Helvetica-Bold',14);c.drawString(58,top-27,title)
 y=top-70
 for f in fields:
  c.setFillColor(ink);c.setFont('Helvetica',11);c.drawString(58,y,f);c.setStrokeColor(HexColor('#aab5aa'));c.line(58,y-29,552,y-29);y-=52
 c.setFont('Helvetica',9);c.drawString(58,top-299,'LIVSOLV.com | Fill in before an emergency. Keep sensitive details protected.')
c.save()
for f in out.glob('*.pdf'):
 pdf=PdfReader(f);assert len(pdf.pages)==(64 if 'Binder' in f.name else 5 if 'Cards' in f.name else 1)
manifest={f.name:hashlib.sha256(f.read_bytes()).hexdigest() for f in sorted(out.iterdir()) if f.is_file()}
(out/'PACKAGE-CONTENTS.txt').write_text('LIVSOLV Systems Down Household Binder\n\nChoose ONE binder print edition. Grayscale versions are black-and-white, not redesigned low-ink layouts. The approved US Letter binder is unchanged.\n\n'+ '\n'.join(manifest)+'\n\nBinder: 64 physical pages, including 62 numbered pages. Quick Cards: 10 cards on 5 sheets. Calculator: 9 tabs.\n')
zip_path=out.parent/'LIVSOLV-Systems-Down-Household-Binder.zip'
with zipfile.ZipFile(zip_path,'w',zipfile.ZIP_DEFLATED) as z:
 for f in sorted(out.iterdir()):
  if f.is_file():z.write(f,f.name)
assert zipfile.ZipFile(zip_path).testzip() is None
print(json.dumps({'zip':str(zip_path),'sha256':hashlib.sha256(zip_path.read_bytes()).hexdigest(),'files':manifest},indent=2))
